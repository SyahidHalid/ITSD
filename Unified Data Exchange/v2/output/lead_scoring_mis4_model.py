# """
# Lead scoring model
# Confirmed customer source: MIS4_EXPANDED.xlsx
# Lead sources: CBRS, ELSA, EPC, JOMXCESS

# Outputs:
# - lead_scoring_mis4_results.xlsx
# - lead_scoring_mis4_results.json
# - lead_scoring_mis4_metrics.json

# Install once:
# pip install pandas numpy openpyxl scikit-learn
# """

import json
import os
import re
import unicodedata
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier
from sklearn.impute import SimpleImputer
from sklearn.metrics import accuracy_score, average_precision_score, f1_score, precision_score, recall_score, roc_auc_score
from sklearn.model_selection import GroupShuffleSplit
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder

# ============================================================
# 1. CONFIGURATION
# ============================================================
BASE_FOLDER = "."

MIS4_FILE = os.path.join(BASE_FOLDER, "MIS4_EXPANDED.xlsx")
CBRS_FILE = os.path.join(BASE_FOLDER, "CBRS_GOLDEN_RECORD_TEST_100.xlsx")
ELSA_FILE = os.path.join(BASE_FOLDER, "ELSA_EXPANDED.xlsx")
EPC_FILE = os.path.join(BASE_FOLDER, "EPC_EXPANDED.xlsx")
JOMXCESS_FILE = os.path.join(BASE_FOLDER, "JOMXCESS_EXPANDED.xlsx")

OUTPUT_EXCEL = os.path.join(BASE_FOLDER, "lead_scoring_mis4_results.xlsx")
OUTPUT_JSON = os.path.join(BASE_FOLDER, "lead_scoring_mis4_results.json")
METRICS_JSON = os.path.join(BASE_FOLDER, "lead_scoring_mis4_metrics.json")

RANDOM_STATE = 42
PROPENSITY_WEIGHT = 0.70
READINESS_WEIGHT = 0.30
VERY_HOT_THRESHOLD = 80
HOT_THRESHOLD = 60
WARM_THRESHOLD = 35

# User-confirmed programme/training sources.
TRAINING_SOURCES = {"CBRS", "ELSA"}

# ============================================================
# 2. HELPERS
# ============================================================
def section(title):
    print("\n" + "=" * 72)
    print(title)
    print("=" * 72)


def validate_files(paths):
    missing = [p for p in paths if not os.path.exists(p)]
    if missing:
        raise FileNotFoundError("Missing files:\n- " + "\n- ".join(missing))


def load_excel(path):
    df = pd.read_excel(path, engine="openpyxl", dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    print(f"Loaded {os.path.basename(path)}: {len(df):,} rows")
    return df


def norm_col(value):
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    return re.sub(r"\s+", " ", value)


def find_col(df, candidates):
    lookup = {norm_col(c): c for c in df.columns}
    for candidate in candidates:
        key = norm_col(candidate)
        if key in lookup:
            return lookup[key]
    for candidate in candidates:
        key = norm_col(candidate)
        if len(key) < 5:
            continue
        for normalised, original in lookup.items():
            if key in normalised or normalised in key:
                return original
    return None


def series(df, candidates, default=""):
    col = find_col(df, candidates)
    if col is None:
        return pd.Series(default, index=df.index)
    return df[col]


def combine(df, candidates, default=""):
    result = pd.Series(default, index=df.index, dtype=object)
    used = []
    for candidate in candidates:
        col = find_col(df, [candidate])
        if col is None or col in used:
            continue
        used.append(col)
        incoming = df[col]
        empty = result.isna() | result.astype(str).str.strip().eq("")
        valid = incoming.notna() & incoming.astype(str).str.strip().ne("")
        result.loc[empty & valid] = incoming.loc[empty & valid]
    return result


def normalise_name(value):
    if pd.isna(value):
        return ""
    text = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii").upper()
    text = re.sub(r"[^A-Z0-9 ]", " ", text)
    for pattern in [
        r"\bSENDIRIAN BERHAD\b", r"\bSDN BHD\b", r"\bSDN BERHAD\b",
        r"\bSDN\b", r"\bBERHAD\b", r"\bBHD\b", r"\bENTERPRISE\b",
        r"\bTRADING\b", r"\bRESOURCES\b", r"\bRESOURCE\b"
    ]:
        text = re.sub(pattern, " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalise_email(value):
    if pd.isna(value):
        return ""
    text = re.sub(r"\s+", "", str(value).lower().strip())
    return text if "@" in text else ""


def normalise_phone(value):
    if pd.isna(value):
        return ""
    text = re.sub(r"\D", "", str(value))
    if text.startswith("60"):
        text = "0" + text[2:]
    return text


def normalise_registration(value):
    if pd.isna(value):
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def as_number(value):
    if pd.isna(value):
        return np.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    text = str(value).upper().replace("RM", "").replace(",", "").replace("%", "").strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else np.nan


def as_yes_no(value):
    if pd.isna(value):
        return np.nan
    text = str(value).upper().strip()
    if text in {"YES", "YA", "Y", "TRUE", "1"} or text.startswith("YES"):
        return 1.0
    if text in {"NO", "TIDAK", "N", "FALSE", "0"} or text.startswith("NO"):
        return 0.0
    return np.nan


def add_keys(df):
    df = df.copy()
    df["MATCH_NAME"] = df["COMPANY_NAME"].map(normalise_name)
    df["MATCH_REGISTRATION"] = df["REGISTRATION_NUMBER"].map(normalise_registration)
    df["MATCH_EMAIL"] = df["EMAIL"].map(normalise_email)
    df["MATCH_PHONE"] = df["PHONE"].map(normalise_phone)
    return df

# ============================================================
# 3. LOAD DATA
# ============================================================
section("LOADING DATA")
validate_files([MIS4_FILE, CBRS_FILE, ELSA_FILE, EPC_FILE, JOMXCESS_FILE])
mis4 = load_excel(MIS4_FILE)
cbrs = load_excel(CBRS_FILE)
elsa = load_excel(ELSA_FILE)
epc = load_excel(EPC_FILE)
jom = load_excel(JOMXCESS_FILE)

# ============================================================
# 4. CONFIRMED CUSTOMER MASTER: MIS4 ONLY
# ============================================================
section("PREPARING MIS4 CUSTOMER MASTER")
customers = pd.DataFrame({
    "CUSTOMER_NAME": combine(mis4, ["CUSTOMER NAME", "COMPANY NAME", "NAMA SYARIKAT"]),
    "REGISTRATION_NUMBER": combine(mis4, ["REGISTRATION NUMBER", "COMPANY REGISTRATION NUMBER", "NO PENDAFTARAN SYARIKAT"]),
    "EMAIL": combine(mis4, ["EMAIL", "EMEL", "EMAIL ADDRESS"]),
    "PHONE": combine(mis4, ["CONTACT NO", "PHONE", "TELEPHONE", "NO TELEFON"]),
})
customers = customers.rename(columns={"CUSTOMER_NAME": "COMPANY_NAME"})
customers = add_keys(customers)
customers = customers.loc[customers[["MATCH_NAME", "MATCH_REGISTRATION", "MATCH_EMAIL", "MATCH_PHONE"]].ne("").any(axis=1)].copy()
customers = customers.drop_duplicates(subset=["MATCH_NAME", "MATCH_REGISTRATION", "MATCH_EMAIL", "MATCH_PHONE"])

customer_names = set(customers.loc[customers.MATCH_NAME.ne(""), "MATCH_NAME"])
customer_regs = set(customers.loc[customers.MATCH_REGISTRATION.ne(""), "MATCH_REGISTRATION"])
customer_emails = set(customers.loc[customers.MATCH_EMAIL.ne(""), "MATCH_EMAIL"])
customer_phones = set(customers.loc[customers.MATCH_PHONE.ne(""), "MATCH_PHONE"])
print(f"MIS4 customer reference records: {len(customers):,}")

# ============================================================
# 5. BUILD FOUR LEAD SOURCES
# ============================================================
section("PREPARING LEAD SOURCES")

cbrs_leads = pd.DataFrame({
    "LEAD_SOURCE": "CBRS",
    "LEAD_DATE": combine(cbrs, ["CREATED_AT", "UPDATED_AT", "TARIKH PENDAFTARAN"]),
    "COMPANY_NAME": combine(cbrs, ["NAMA SYARIKAT", "COMPANY NAME"]),
    "REGISTRATION_NUMBER": combine(cbrs, ["NO PENDAFTARAN SYARIKAT YANG DIDAFTARKAN (SSM/PBT)", "REGISTRATION NUMBER"]),
    "EMAIL": combine(cbrs, ["EMAIL", "EMEL"]),
    "PHONE": combine(cbrs, ["NO TELEFON SYARIKAT", "NO_TELEFON_BIMBIT", "PHONE"]),
    "STATE": combine(cbrs, ["NEGERI", "STATE"]),
    "SECTOR": combine(cbrs, ["KATEGORI SEKTOR", "AKTIVITI PERNIAGAAN YANG DIJALANKAN", "SECTOR"]),
    "SME_CATEGORY": combine(cbrs, ["SAIZ PERNIAGAAN", "SME CATEGORY"]),
    "PROGRAMME": combine(cbrs, ["PROGRAMME", "PROGRAM"]),
    "REFERRAL_SOURCE": combine(cbrs, ["SUMBER RUJUKAN", "REFERRAL SOURCE"]),
    "ANNUAL_TURNOVER": np.nan,
    "FULL_TIME_EMPLOYEES": np.nan,
    "ATTENDANCE": np.nan,
    "COLLECTION_RATE": np.nan,
    "AGEING": np.nan,
    "ECL": np.nan,
    "BUSINESS_NEED": "",
    "HAS_BANK_ACCOUNT": np.nan,
    "HAS_AUDITED_ACCOUNT": np.nan,
    "SST_REGISTERED": np.nan,
    "HAS_BUSINESS_PLAN": np.nan,
    "ATTEMPTED_FINANCING": np.nan,
    "SOURCE_SCORE": np.nan,
})

elsa_leads = pd.DataFrame({
    "LEAD_SOURCE": "ELSA",
    "LEAD_DATE": combine(elsa, ["ENTRY_DATE", "TIMESTAMP", "DATE"]),
    "COMPANY_NAME": combine(elsa, ["Nama Syarikat", "Company Name"]),
    "REGISTRATION_NUMBER": combine(elsa, ["Nombor Pendaftaran Baru", "Nombor Pendaftaran Lama", "Registration Number"]),
    "EMAIL": combine(elsa, ["Emel", "Email"]),
    "PHONE": combine(elsa, ["Nombor telefon", "No Telefon", "Contact No"]),
    "STATE": combine(elsa, ["Negeri", "State"]),
    "SECTOR": combine(elsa, ["Sektor", "Sector", "Industry"]),
    "SME_CATEGORY": combine(elsa, ["Kategori SME", "SME Category"]),
    "PROGRAMME": "ELSA",
    "REFERRAL_SOURCE": combine(elsa, ["Sumber Rujukan", "Referral Source"]),
    "ANNUAL_TURNOVER": combine(elsa, ["Perolehan Tahunan (RM)", "Perolehan Tahunan", "Annual Turnover"], np.nan).map(as_number),
    "FULL_TIME_EMPLOYEES": combine(elsa, ["Berapakah jumlah pekerja yang bekerja secara sepenuh masa? (Sila nyatakan angka)", "Full Time Employees"], np.nan).map(as_number),
    "ATTENDANCE": np.nan,
    "COLLECTION_RATE": np.nan,
    "AGEING": np.nan,
    "ECL": np.nan,
    "BUSINESS_NEED": combine(elsa, ["Keperluan", "Business Need"]),
    "HAS_BANK_ACCOUNT": combine(elsa, ["Adakah anda mempunyai akaun bank berdaftar untuk syarikat anda?"], np.nan).map(as_yes_no),
    "HAS_AUDITED_ACCOUNT": combine(elsa, ["Adakah anda mempunyai akaun perniagaan yang diaudit?"], np.nan).map(as_yes_no),
    "SST_REGISTERED": combine(elsa, ["Adakah anda berdaftar dengan SST?"], np.nan).map(as_yes_no),
    "HAS_BUSINESS_PLAN": combine(elsa, ["Adakah anda mempunyai Perancangan Perniagaan yang didokumentasikan untuk syarikat anda?"], np.nan).map(as_yes_no),
    "ATTEMPTED_FINANCING": combine(elsa, ["Adakah anda pernah mencuba mendapatkan pembiayaan untuk perniagaan anda sejak dua tahun kebelakangan ini?"], np.nan).map(as_yes_no),
    "SOURCE_SCORE": combine(elsa, ["SCORE", "TOTAL SCORE", "ELSA SCORE"], np.nan).map(as_number),
})

epc_leads = pd.DataFrame({
    "LEAD_SOURCE": "EPC",
    "LEAD_DATE": "",
    "COMPANY_NAME": combine(epc, ["PENYEWA", "COMPANY NAME", "CUSTOMER NAME"]),
    "REGISTRATION_NUMBER": combine(epc, ["REGISTRATION NUMBER", "NO PENDAFTARAN SYARIKAT"]),
    "EMAIL": combine(epc, ["EMAIL", "EMEL"]),
    "PHONE": combine(epc, ["CONTACT NO", "PHONE", "NO TELEFON"]),
    "STATE": combine(epc, ["STATE", "NEGERI"]),
    "SECTOR": combine(epc, ["INDUSTRY", "SECTOR"]),
    "SME_CATEGORY": "",
    "PROGRAMME": "EPC",
    "REFERRAL_SOURCE": "",
    "ANNUAL_TURNOVER": np.nan,
    "FULL_TIME_EMPLOYEES": np.nan,
    "ATTENDANCE": np.nan,
    "COLLECTION_RATE": combine(epc, ["% TOTAL COLLECTION/DUE JUNE 2026", "COLLECTION RATE"], np.nan).map(as_number),
    "AGEING": combine(epc, ["AGEING JUNE 2026", "AGEING", "AGING"], np.nan).map(as_number),
    "ECL": combine(epc, ["ECL PROVISION FOR DOUBTFUL JUNE 2026", "ECL"], np.nan).map(as_number),
    "BUSINESS_NEED": "",
    "HAS_BANK_ACCOUNT": np.nan,
    "HAS_AUDITED_ACCOUNT": np.nan,
    "SST_REGISTERED": np.nan,
    "HAS_BUSINESS_PLAN": np.nan,
    "ATTEMPTED_FINANCING": np.nan,
    "SOURCE_SCORE": np.nan,
})

jom_leads = pd.DataFrame({
    "LEAD_SOURCE": "JOMXCESS",
    "LEAD_DATE": combine(jom, ["Timestamp", "Date"]),
    "COMPANY_NAME": combine(jom, ["Nama Syarikat", "Nama Syarikat 2", "Company Name"]),
    "REGISTRATION_NUMBER": combine(jom, ["Registration Number", "No Pendaftaran Syarikat"]),
    "EMAIL": combine(jom, ["Email", "Emel"]),
    "PHONE": combine(jom, ["No telefon / pejabat", "No. telefon bimbit / pejabat", "Contact No"]),
    "STATE": combine(jom, ["Negeri", "State"]),
    "SECTOR": combine(jom, ["Jenis Perniagaan (Nature of Business)?", "Nature of Business", "Sector"]),
    "SME_CATEGORY": combine(jom, ["Kategori ", "Kategori  2", "SME Category"]),
    "PROGRAMME": "JOMXCESS",
    "REFERRAL_SOURCE": combine(jom, ["Sumber Rujukan", "Referral Source"]),
    "ANNUAL_TURNOVER": np.nan,
    "FULL_TIME_EMPLOYEES": np.nan,
    "ATTENDANCE": combine(jom, ["Jumlah kehadiran ", "Jumlah kehadiran 2", "Attendance"], np.nan).map(as_number),
    "COLLECTION_RATE": np.nan,
    "AGEING": np.nan,
    "ECL": np.nan,
    "BUSINESS_NEED": combine(jom, ["Keperluan anda? ", "Business Need"]),
    "HAS_BANK_ACCOUNT": np.nan,
    "HAS_AUDITED_ACCOUNT": np.nan,
    "SST_REGISTERED": np.nan,
    "HAS_BUSINESS_PLAN": np.nan,
    "ATTEMPTED_FINANCING": np.nan,
    "SOURCE_SCORE": np.nan,
})

leads = pd.concat([cbrs_leads, elsa_leads, epc_leads, jom_leads], ignore_index=True, sort=False)
leads.insert(0, "LEAD_ID", [f"LEAD-{i:06d}" for i in range(1, len(leads) + 1)])
leads["TRAINING_PARTICIPANT"] = leads["LEAD_SOURCE"].isin(TRAINING_SOURCES).astype(int)
leads = add_keys(leads)
leads = leads.loc[leads[["MATCH_NAME", "MATCH_REGISTRATION", "MATCH_EMAIL", "MATCH_PHONE"]].ne("").any(axis=1)].reset_index(drop=True)
print(leads.groupby("LEAD_SOURCE").size())

# ============================================================
# 6. MATCH LEADS AGAINST MIS4
# ============================================================
section("MATCHING LEADS TO MIS4")
leads["MATCHED_BY_REGISTRATION"] = leads.MATCH_REGISTRATION.ne("") & leads.MATCH_REGISTRATION.isin(customer_regs)
leads["MATCHED_BY_EMAIL"] = leads.MATCH_EMAIL.ne("") & leads.MATCH_EMAIL.isin(customer_emails)
leads["MATCHED_BY_PHONE"] = leads.MATCH_PHONE.ne("") & leads.MATCH_PHONE.isin(customer_phones)
leads["MATCHED_BY_NAME"] = leads.MATCH_NAME.ne("") & leads.MATCH_NAME.isin(customer_names)
match_flags = ["MATCHED_BY_REGISTRATION", "MATCHED_BY_EMAIL", "MATCHED_BY_PHONE", "MATCHED_BY_NAME"]
leads["MATCH_COUNT"] = leads[match_flags].astype(int).sum(axis=1)
leads["IS_MIS4_CUSTOMER"] = (leads.MATCH_COUNT > 0).astype(int)


def match_method(row):
    labels = []
    for col, label in [
        ("MATCHED_BY_REGISTRATION", "REGISTRATION"),
        ("MATCHED_BY_EMAIL", "EMAIL"),
        ("MATCHED_BY_PHONE", "PHONE"),
        ("MATCHED_BY_NAME", "COMPANY_NAME"),
    ]:
        if row[col]:
            labels.append(label)
    return ", ".join(labels) if labels else "NO_MATCH"


leads["MATCH_METHOD"] = leads.apply(match_method, axis=1)

# ============================================================
# 7. READINESS SCORE
# ============================================================
section("CALCULATING READINESS")


def readiness(row):
    earned, possible = 0.0, 0.0

    def binary(value, weight):
        nonlocal earned, possible
        if pd.notna(value):
            possible += weight
            earned += weight if float(value) == 1 else 0

    def scaled(value, weight, upper):
        nonlocal earned, possible
        if pd.notna(value):
            possible += weight
            earned += min(max(float(value), 0) / upper, 1) * weight

    # Training participation is intentionally important because CBRS and ELSA
    # participants have already received financing-related training.
    binary(row.get("TRAINING_PARTICIPANT"), 20)
    binary(row.get("HAS_BANK_ACCOUNT"), 8)
    binary(row.get("HAS_AUDITED_ACCOUNT"), 10)
    binary(row.get("SST_REGISTERED"), 6)
    binary(row.get("HAS_BUSINESS_PLAN"), 8)
    binary(row.get("ATTEMPTED_FINANCING"), 10)

    turnover = row.get("ANNUAL_TURNOVER")
    if pd.notna(turnover):
        possible += 15
        earned += min(np.log1p(max(float(turnover), 0)) / np.log1p(50_000_000), 1) * 15

    scaled(row.get("FULL_TIME_EMPLOYEES"), 7, 75)
    scaled(row.get("ATTENDANCE"), 10, 5)

    score = row.get("SOURCE_SCORE")
    if pd.notna(score):
        possible += 8
        score = float(score)
        earned += min(max(score / (5 if score <= 5 else 100), 0), 1) * 8

    collection = row.get("COLLECTION_RATE")
    if pd.notna(collection):
        possible += 15
        collection = float(collection)
        if collection <= 1.5:
            collection *= 100
        earned += min(max(collection, 0), 100) / 100 * 15

    ageing = row.get("AGEING")
    if pd.notna(ageing):
        possible += 10
        earned += max(0, 1 - min(max(float(ageing), 0), 180) / 180) * 10

    need = str(row.get("BUSINESS_NEED", "")).upper()
    if need.strip():
        possible += 10
        if any(x in need for x in ["PEMBIAYAAN", "FINANC", "LOAN"]):
            earned += 10
        elif any(x in need for x in ["PEMBANGUNAN", "DEVELOPMENT", "CAPACITY"]):
            earned += 7
        else:
            earned += 4

    return round(50.0 if possible == 0 else 100 * earned / possible, 2)


leads["READINESS_SCORE"] = leads.apply(readiness, axis=1)

# ============================================================
# 8. PROPENSITY MODEL
# ============================================================
section("TRAINING PROPENSITY MODEL")
feature_cols = [
    "LEAD_SOURCE", "TRAINING_PARTICIPANT", "STATE", "SECTOR", "SME_CATEGORY",
    "PROGRAMME", "REFERRAL_SOURCE", "ANNUAL_TURNOVER", "FULL_TIME_EMPLOYEES",
    "ATTENDANCE", "COLLECTION_RATE", "AGEING", "ECL", "BUSINESS_NEED",
    "HAS_BANK_ACCOUNT", "HAS_AUDITED_ACCOUNT", "SST_REGISTERED",
    "HAS_BUSINESS_PLAN", "ATTEMPTED_FINANCING", "SOURCE_SCORE", "READINESS_SCORE"
]
X = leads[feature_cols].copy()
y = leads["IS_MIS4_CUSTOMER"].copy()
cat_cols = [c for c in feature_cols if X[c].dtype == "object"]
num_cols = [c for c in feature_cols if c not in cat_cols]

# Ensure categorical columns contain one consistent data type.
for column in cat_cols:
    X[column] = X[column].fillna("MISSING").astype(str).replace("", "MISSING")

for column in num_cols:
    X[column] = pd.to_numeric(X[column], errors="coerce")

preprocessor = ColumnTransformer([
    ("categorical", Pipeline([
        ("imputer", SimpleImputer(strategy="most_frequent")),
        ("encoder", OneHotEncoder(handle_unknown="ignore")),
    ]), cat_cols),
    ("numeric", Pipeline([
        ("imputer", SimpleImputer(strategy="median")),
    ]), num_cols),
])

model = Pipeline([
    ("preprocessor", preprocessor),
    ("classifier", RandomForestClassifier(
        n_estimators=500,
        min_samples_leaf=2,
        class_weight="balanced",
        random_state=RANDOM_STATE,
        n_jobs=-1,
    )),
])

metrics = {
    "generated_at": datetime.now().isoformat(),
    "confirmed_customer_source": "MIS4",
    "lead_sources": ["CBRS", "ELSA", "EPC", "JOMXCESS"],
    "training_sources": sorted(TRAINING_SOURCES),
    "total_records": int(len(leads)),
    "matched_to_mis4": int(y.sum()),
    "not_matched_to_mis4": int((y == 0).sum()),
    "model_trained": False,
    "warning": None,
}

can_train = y.nunique() == 2 and y.value_counts().min() >= 3 and len(y) >= 20
if can_train:
    groups = leads.MATCH_NAME.replace("", np.nan)
    groups = groups.fillna(pd.Series([f"ROW-{i}" for i in leads.index], index=leads.index))
    splitter = GroupShuffleSplit(n_splits=1, test_size=0.20, random_state=RANDOM_STATE)
    train_idx, test_idx = next(splitter.split(X, y, groups=groups))
    X_train, X_test = X.iloc[train_idx], X.iloc[test_idx]
    y_train, y_test = y.iloc[train_idx], y.iloc[test_idx]

    if y_train.nunique() == 2:
        model.fit(X_train, y_train)
        leads["CONVERSION_PROPENSITY"] = model.predict_proba(X)[:, 1] * 100
        test_probability = model.predict_proba(X_test)[:, 1]
        test_prediction = (test_probability >= 0.50).astype(int)
        metrics.update({
            "model_trained": True,
            "training_records": int(len(train_idx)),
            "testing_records": int(len(test_idx)),
            "accuracy": float(accuracy_score(y_test, test_prediction)),
            "precision": float(precision_score(y_test, test_prediction, zero_division=0)),
            "recall": float(recall_score(y_test, test_prediction, zero_division=0)),
            "f1_score": float(f1_score(y_test, test_prediction, zero_division=0)),
            "roc_auc": float(roc_auc_score(y_test, test_probability)) if y_test.nunique() == 2 else None,
            "pr_auc": float(average_precision_score(y_test, test_probability)) if y_test.nunique() == 2 else None,
        })
    else:
        can_train = False

if not can_train or not metrics["model_trained"]:
    leads["CONVERSION_PROPENSITY"] = leads["READINESS_SCORE"]
    metrics["warning"] = "Insufficient class variation after group split; readiness score used as provisional propensity."

# ============================================================
# 9. FINAL RANKING
# ============================================================
section("RANKING NEW LEADS")
leads["CONVERSION_PROPENSITY"] = leads["CONVERSION_PROPENSITY"].clip(0, 100).round(2)
leads["FINAL_LEAD_SCORE"] = (
    PROPENSITY_WEIGHT * leads["CONVERSION_PROPENSITY"]
    + READINESS_WEIGHT * leads["READINESS_SCORE"]
).round(2)


def priority(row):
    if row.IS_MIS4_CUSTOMER == 1:
        return "EXISTING MIS4 CUSTOMER"
    if row.FINAL_LEAD_SCORE >= VERY_HOT_THRESHOLD:
        return "VERY HOT"
    if row.FINAL_LEAD_SCORE >= HOT_THRESHOLD:
        return "HOT"
    if row.FINAL_LEAD_SCORE >= WARM_THRESHOLD:
        return "WARM"
    return "COLD"


leads["LEAD_PRIORITY"] = leads.apply(priority, axis=1)
leads["RECOMMENDED_ACTION"] = leads.LEAD_PRIORITY.map({
    "VERY HOT": "Immediate relationship manager review",
    "HOT": "Contact for financing qualification",
    "WARM": "Nurture and collect additional information",
    "COLD": "Retain in lead pool for future engagement",
    "EXISTING MIS4 CUSTOMER": "Exclude from new-lead campaign",
})

leads["DEDUP_KEY"] = np.where(
    leads.MATCH_REGISTRATION.ne(""), "REG|" + leads.MATCH_REGISTRATION,
    np.where(leads.MATCH_EMAIL.ne(""), "EMAIL|" + leads.MATCH_EMAIL,
             np.where(leads.MATCH_PHONE.ne(""), "PHONE|" + leads.MATCH_PHONE,
                      "NAME|" + leads.MATCH_NAME))
)

ranked = leads.sort_values(["IS_MIS4_CUSTOMER", "FINAL_LEAD_SCORE"], ascending=[True, False])
ranked = ranked.drop_duplicates(["LEAD_SOURCE", "DEDUP_KEY"], keep="first")
new_leads = ranked.loc[ranked.IS_MIS4_CUSTOMER.eq(0)].sort_values("FINAL_LEAD_SCORE", ascending=False).reset_index(drop=True)
existing = ranked.loc[ranked.IS_MIS4_CUSTOMER.eq(1)].sort_values(["LEAD_SOURCE", "COMPANY_NAME"]).reset_index(drop=True)
new_leads.insert(0, "LEAD_RANK", range(1, len(new_leads) + 1))

output_cols = [
    "LEAD_RANK", "LEAD_ID", "LEAD_SOURCE", "TRAINING_PARTICIPANT", "PROGRAMME",
    "COMPANY_NAME", "REGISTRATION_NUMBER", "EMAIL", "PHONE", "STATE", "SECTOR",
    "SME_CATEGORY", "BUSINESS_NEED", "ANNUAL_TURNOVER", "FULL_TIME_EMPLOYEES",
    "ATTENDANCE", "COLLECTION_RATE", "AGEING", "READINESS_SCORE",
    "CONVERSION_PROPENSITY", "FINAL_LEAD_SCORE", "LEAD_PRIORITY", "RECOMMENDED_ACTION"
]
new_output = new_leads[[c for c in output_cols if c in new_leads.columns]].copy()
existing_output = existing[[
    "LEAD_ID", "LEAD_SOURCE", "COMPANY_NAME", "REGISTRATION_NUMBER", "EMAIL", "PHONE",
    "MATCH_METHOD", "MATCH_COUNT", "READINESS_SCORE", "CONVERSION_PROPENSITY",
    "FINAL_LEAD_SCORE", "LEAD_PRIORITY", "RECOMMENDED_ACTION"
]].copy()

source_summary = new_output.groupby("LEAD_SOURCE", dropna=False).agg(
    TOTAL_LEADS=("LEAD_ID", "count"),
    TRAINED_LEADS=("TRAINING_PARTICIPANT", "sum"),
    AVERAGE_READINESS=("READINESS_SCORE", "mean"),
    AVERAGE_PROPENSITY=("CONVERSION_PROPENSITY", "mean"),
    AVERAGE_FINAL_SCORE=("FINAL_LEAD_SCORE", "mean"),
).reset_index().round(2)

priority_summary = new_output.groupby(["LEAD_SOURCE", "LEAD_PRIORITY"], dropna=False).size().reset_index(name="TOTAL_LEADS")
matching_details = leads[[
    "LEAD_ID", "LEAD_SOURCE", "COMPANY_NAME", "MATCH_NAME", "MATCH_REGISTRATION",
    "MATCH_EMAIL", "MATCH_PHONE", *match_flags, "MATCH_COUNT", "MATCH_METHOD", "IS_MIS4_CUSTOMER"
]].copy()

# ============================================================
# 10. EXPORT EXCEL
# ============================================================
section("EXPORTING EXCEL AND JSON")
with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
    new_output.to_excel(writer, sheet_name="New Lead Rankings", index=False)
    existing_output.to_excel(writer, sheet_name="Existing MIS4 Matches", index=False)
    source_summary.to_excel(writer, sheet_name="Source Summary", index=False)
    priority_summary.to_excel(writer, sheet_name="Priority Summary", index=False)
    matching_details.to_excel(writer, sheet_name="Matching Details", index=False)

workbook = load_workbook(OUTPUT_EXCEL)
header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
priority_colours = {
    "VERY HOT": "C6EFCE", "HOT": "FFEB9C", "WARM": "FCE4D6",
    "COLD": "E7E6E6", "EXISTING MIS4 CUSTOMER": "D9EAF7"
}

for worksheet in workbook.worksheets:
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    # Correct syntax: format every cell in the first row.
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column_cells in worksheet.columns:
        max_length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 45)

    header_map = {cell.value: cell.column for cell in worksheet[1]}
    priority_col = header_map.get("LEAD_PRIORITY")
    if priority_col:
        for row_number in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_number, column=priority_col)
            if cell.value in priority_colours:
                cell.fill = PatternFill("solid", fgColor=priority_colours[cell.value])
                cell.font = Font(bold=True)

workbook.save(OUTPUT_EXCEL)

# JSON contains only new, ranked prospects.
json_df = new_output.replace({np.nan: None, np.inf: None, -np.inf: None})
with open(OUTPUT_JSON, "w", encoding="utf-8") as file:
    json.dump(json_df.to_dict(orient="records"), file, ensure_ascii=False, indent=2, default=str)

with open(METRICS_JSON, "w", encoding="utf-8") as file:
    json.dump(metrics, file, ensure_ascii=False, indent=2, default=str)

print(f"Created: {OUTPUT_EXCEL}")
print(f"Created: {OUTPUT_JSON}")
print(f"Created: {METRICS_JSON}")
print(f"New unique leads ranked: {len(new_output):,}")
print(f"Lead records matching MIS4: {int(leads.IS_MIS4_CUSTOMER.sum()):,}")
print("\nNote: Propensity is a prioritisation score. For a validated conversion model, use historical lead dates and later MIS4 onboarding outcomes.")
