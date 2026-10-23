# ============================================================
# CUSTOMER LEAD SCORING MODEL
#
# Existing customer sources:
#   - CBRS_GOLDEN_RECORD_TEST_100.xlsx
#   - MIS4_EXPANDED.xlsx
#
# Potential lead sources:
#   - ELSA_EXPANDED.xlsx
#   - EPC_EXPANDED.xlsx
#   - JOMXCESS_EXPANDED.xlsx
#
# Outputs:
#   - lead_scoring_results.xlsx
#   - lead_scoring_results.json
#   - lead_scoring_metrics.json
#   - lead_propensity_model.cbm, if model training succeeds
#
# First-time installation:
# pip install pandas numpy openpyxl scikit-learn catboost
# ============================================================

import os
import re
import json
import warnings
import unicodedata
from datetime import datetime

import numpy as np
import pandas as pd

from sklearn.model_selection import GroupShuffleSplit
from sklearn.metrics import (
    accuracy_score,
    precision_score,
    recall_score,
    f1_score,
    roc_auc_score,
    average_precision_score,
    confusion_matrix
)

from catboost import CatBoostClassifier

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


warnings.filterwarnings("ignore")


# ============================================================
# 1. CONFIGURATION
# ============================================================

# Use "." when the script/notebook is in the same folder
# as the Excel files.
BASE_FOLDER = "."

CBRS_FILE = os.path.join(
    BASE_FOLDER,
    "CBRS_GOLDEN_RECORD_TEST_100.xlsx"
)

MIS4_FILE = os.path.join(
    BASE_FOLDER,
    "MIS4_EXPANDED.xlsx"
)

ELSA_FILE = os.path.join(
    BASE_FOLDER,
    "ELSA_EXPANDED.xlsx"
)

EPC_FILE = os.path.join(
    BASE_FOLDER,
    "EPC_EXPANDED.xlsx"
)

JOMXCESS_FILE = os.path.join(
    BASE_FOLDER,
    "JOMXCESS_EXPANDED.xlsx"
)

OUTPUT_EXCEL = os.path.join(
    BASE_FOLDER,
    "lead_scoring_results.xlsx"
)

OUTPUT_JSON = os.path.join(
    BASE_FOLDER,
    "lead_scoring_results.json"
)

METRICS_JSON = os.path.join(
    BASE_FOLDER,
    "lead_scoring_metrics.json"
)

MODEL_FILE = os.path.join(
    BASE_FOLDER,
    "lead_propensity_model.cbm"
)

RANDOM_STATE = 42

PROPENSITY_WEIGHT = 0.70
READINESS_WEIGHT = 0.30

VERY_HOT_THRESHOLD = 80
HOT_THRESHOLD = 60
WARM_THRESHOLD = 35


# ============================================================
# 2. GENERAL FUNCTIONS
# ============================================================

def print_section(title):
    print("\n" + "=" * 70)
    print(title)
    print("=" * 70)


def validate_files(file_paths):
    missing_files = [
        path
        for path in file_paths
        if not os.path.exists(path)
    ]

    if missing_files:
        raise FileNotFoundError(
            "The following files were not found:\n\n"
            + "\n".join(missing_files)
            + "\n\nCheck BASE_FOLDER and the filenames."
        )


def load_excel(file_path):
    print(f"Loading: {file_path}")

    dataframe = pd.read_excel(
        file_path,
        engine="openpyxl",
        dtype=object
    )

    dataframe.columns = [
        str(column).strip()
        for column in dataframe.columns
    ]

    print(
        f"Loaded {len(dataframe):,} rows and "
        f"{len(dataframe.columns):,} columns."
    )

    return dataframe


def normalise_column_name(value):
    value = str(value).strip().lower()
    value = re.sub(r"\s+", " ", value)
    value = re.sub(r"[^a-z0-9 ]", "", value)
    return value


def find_column(dataframe, possible_names):
    """
    Finds a column by:
    1. Exact normalised name
    2. Partial normalised name
    """

    lookup = {
        normalise_column_name(column): column
        for column in dataframe.columns
    }

    # Exact matching
    for name in possible_names:
        normalised_name = normalise_column_name(name)

        if normalised_name in lookup:
            return lookup[normalised_name]

    # Partial matching
    for name in possible_names:
        normalised_name = normalise_column_name(name)

        if len(normalised_name) < 4:
            continue

        for normalised_column, original_column in lookup.items():
            if (
                normalised_name in normalised_column
                or normalised_column in normalised_name
            ):
                return original_column

    return None


def get_series(
    dataframe,
    possible_names,
    default_value=""
):
    column = find_column(
        dataframe,
        possible_names
    )

    if column is None:
        return pd.Series(
            default_value,
            index=dataframe.index
        )

    return dataframe[column]


def clean_unicode(value):
    if pd.isna(value):
        return ""

    value = unicodedata.normalize(
        "NFKD",
        str(value)
    )

    return (
        value.encode("ascii", "ignore")
        .decode("ascii")
    )


def normalise_company_name(value):
    if pd.isna(value):
        return ""

    value = clean_unicode(value).upper().strip()

    value = re.sub(
        r"[^A-Z0-9 ]",
        " ",
        value
    )

    company_terms = [
        r"\bSENDIRIAN BERHAD\b",
        r"\bSDN BHD\b",
        r"\bSDN BERHAD\b",
        r"\bSDN\b",
        r"\bBERHAD\b",
        r"\bBHD\b",
        r"\bENTERPRISE\b",
        r"\bENT\b",
        r"\bTRADING\b",
        r"\bRESOURCES\b",
        r"\bRESOURCE\b"
    ]

    for pattern in company_terms:
        value = re.sub(
            pattern,
            " ",
            value
        )

    value = re.sub(
        r"\s+",
        " ",
        value
    ).strip()

    return value


def normalise_email(value):
    if pd.isna(value):
        return ""

    value = str(value).lower().strip()
    value = re.sub(r"\s+", "", value)

    if "@" not in value:
        return ""

    return value


def normalise_phone(value):
    if pd.isna(value):
        return ""

    digits = re.sub(
        r"\D",
        "",
        str(value)
    )

    if not digits:
        return ""

    if digits.startswith("60"):
        digits = "0" + digits[2:]

    return digits


def normalise_registration(value):
    if pd.isna(value):
        return ""

    return re.sub(
        r"[^A-Z0-9]",
        "",
        str(value).upper()
    )


def convert_numeric(value):
    if pd.isna(value):
        return np.nan

    if isinstance(
        value,
        (int, float, np.integer, np.floating)
    ):
        return float(value)

    text = str(value).strip().upper()

    if not text:
        return np.nan

    text = text.replace("RM", "")
    text = text.replace(",", "")
    text = text.replace("%", "")
    text = text.strip()

    try:
        return float(text)
    except (ValueError, TypeError):
        return np.nan


def convert_yes_no(value):
    if pd.isna(value):
        return np.nan

    text = str(value).strip().upper()

    yes_values = {
        "YES",
        "YA",
        "Y",
        "TRUE",
        "1",
        "YES / YA"
    }

    no_values = {
        "NO",
        "TIDAK",
        "N",
        "FALSE",
        "0",
        "NO / TIDAK"
    }

    if text in yes_values:
        return 1.0

    if text in no_values:
        return 0.0

    if text.startswith("YES") or text == "YA":
        return 1.0

    if text.startswith("NO") or text == "TIDAK":
        return 0.0

    return np.nan


def combine_first_available(
    dataframe,
    possible_names,
    default_value=""
):
    """
    Combines several possible columns by taking the first
    non-empty value in each row.
    """

    matching_columns = []

    for name in possible_names:
        column = find_column(
            dataframe,
            [name]
        )

        if (
            column is not None
            and column not in matching_columns
        ):
            matching_columns.append(column)

    if not matching_columns:
        return pd.Series(
            default_value,
            index=dataframe.index
        )

    result = pd.Series(
        default_value,
        index=dataframe.index,
        dtype=object
    )

    for column in matching_columns:
        values = dataframe[column]

        empty_result = (
            result.isna()
            | result.astype(str).str.strip().eq("")
        )

        usable_value = (
            values.notna()
            & values.astype(str).str.strip().ne("")
        )

        mask = empty_result & usable_value

        result.loc[mask] = values.loc[mask]

    return result


# ============================================================
# 3. VALIDATE AND LOAD FILES
# ============================================================

print_section("VALIDATING INPUT FILES")

validate_files([
    CBRS_FILE,
    MIS4_FILE,
    ELSA_FILE,
    EPC_FILE,
    JOMXCESS_FILE
])

print_section("LOADING EXCEL FILES")

cbrs = load_excel(CBRS_FILE)
mis4 = load_excel(MIS4_FILE)
elsa = load_excel(ELSA_FILE)
epc = load_excel(EPC_FILE)
jomxcess = load_excel(JOMXCESS_FILE)


# ============================================================
# 4. DISPLAY IDENTIFIED COLUMNS
# ============================================================

print_section("AVAILABLE COLUMNS")

print("\nCBRS columns:")
print(cbrs.columns.tolist())

print("\nMIS4 columns:")
print(mis4.columns.tolist())

print("\nELSA columns:")
print(elsa.columns.tolist())

print("\nEPC columns:")
print(epc.columns.tolist())

print("\nJOMXCESS columns:")
print(jomxcess.columns.tolist())


# ============================================================
# 5. PREPARE CBRS CUSTOMER MASTER
# ============================================================

print_section("PREPARING CUSTOMER MASTER")

cbrs_customers = pd.DataFrame({
    "CUSTOMER_SOURCE": "CBRS",

    "CUSTOMER_NAME": combine_first_available(
        cbrs,
        [
            "NAMA SYARIKAT",
            "COMPANY NAME",
            "CUSTOMER NAME",
            "NAMA PERNIAGAAN"
        ]
    ),

    "CUSTOMER_REGISTRATION": combine_first_available(
        cbrs,
        [
            "NO PENDAFTARAN SYARIKAT YANG DIDAFTARKAN SSM PBT",
            "NO PENDAFTARAN SYARIKAT",
            "NOMBOR PENDAFTARAN",
            "REGISTRATION NUMBER",
            "COMPANY REGISTRATION NUMBER"
        ]
    ),

    "CUSTOMER_EMAIL": combine_first_available(
        cbrs,
        [
            "EMAIL",
            "EMEL",
            "EMAIL ADDRESS"
        ]
    ),

    "CUSTOMER_PHONE": combine_first_available(
        cbrs,
        [
            "NO TELEFON SYARIKAT",
            "NO TELEFON",
            "CONTACT NO",
            "PHONE",
            "TELEPHONE"
        ]
    )
})


# ============================================================
# 6. PREPARE MIS4 CUSTOMER MASTER
# ============================================================

mis4_customers = pd.DataFrame({
    "CUSTOMER_SOURCE": "MIS4",

    "CUSTOMER_NAME": combine_first_available(
        mis4,
        [
            "CUSTOMER NAME",
            "COMPANY NAME",
            "NAMA SYARIKAT",
            "CUSTOMER_NAME"
        ]
    ),

    "CUSTOMER_REGISTRATION": combine_first_available(
        mis4,
        [
            "COMPANY REGISTRATION NUMBER",
            "REGISTRATION NUMBER",
            "NO PENDAFTARAN SYARIKAT"
        ]
    ),

    "CUSTOMER_EMAIL": combine_first_available(
        mis4,
        [
            "EMAIL",
            "EMEL",
            "EMAIL ADDRESS"
        ]
    ),

    "CUSTOMER_PHONE": combine_first_available(
        mis4,
        [
            "CONTACT NO",
            "PHONE",
            "TELEPHONE",
            "NO TELEFON"
        ]
    )
})


customer_master = pd.concat(
    [
        cbrs_customers,
        mis4_customers
    ],
    ignore_index=True,
    sort=False
)


customer_master["MATCH_NAME"] = (
    customer_master["CUSTOMER_NAME"]
    .map(normalise_company_name)
)

customer_master["MATCH_REGISTRATION"] = (
    customer_master["CUSTOMER_REGISTRATION"]
    .map(normalise_registration)
)

customer_master["MATCH_EMAIL"] = (
    customer_master["CUSTOMER_EMAIL"]
    .map(normalise_email)
)

customer_master["MATCH_PHONE"] = (
    customer_master["CUSTOMER_PHONE"]
    .map(normalise_phone)
)


# Remove records without any usable identifier
customer_master = customer_master.loc[
    customer_master[
        [
            "MATCH_NAME",
            "MATCH_REGISTRATION",
            "MATCH_EMAIL",
            "MATCH_PHONE"
        ]
    ].ne("").any(axis=1)
].copy()


customer_master = customer_master.drop_duplicates(
    subset=[
        "CUSTOMER_SOURCE",
        "MATCH_NAME",
        "MATCH_REGISTRATION",
        "MATCH_EMAIL",
        "MATCH_PHONE"
    ]
).reset_index(drop=True)


customer_names = set(
    customer_master.loc[
        customer_master["MATCH_NAME"].ne(""),
        "MATCH_NAME"
    ]
)

customer_registrations = set(
    customer_master.loc[
        customer_master[
            "MATCH_REGISTRATION"
        ].ne(""),
        "MATCH_REGISTRATION"
    ]
)

customer_emails = set(
    customer_master.loc[
        customer_master["MATCH_EMAIL"].ne(""),
        "MATCH_EMAIL"
    ]
)

customer_phones = set(
    customer_master.loc[
        customer_master["MATCH_PHONE"].ne(""),
        "MATCH_PHONE"
    ]
)


print(
    f"Customer master records: "
    f"{len(customer_master):,}"
)


# ============================================================
# 7. PREPARE ELSA LEADS
# ============================================================

print_section("PREPARING ELSA LEADS")

elsa_leads = pd.DataFrame({
    "LEAD_SOURCE": "ELSA",

    "LEAD_DATE": combine_first_available(
        elsa,
        [
            "ENTRY_DATE",
            "TIMESTAMP",
            "DATE"
        ]
    ),

    "COMPANY_NAME": combine_first_available(
        elsa,
        [
            "Nama Syarikat",
            "Company Name",
            "Nama Perniagaan"
        ]
    ),

    "REGISTRATION_NUMBER": combine_first_available(
        elsa,
        [
            "Nombor Pendaftaran Baru",
            "Nombor Pendaftaran Lama",
            "No Pendaftaran Syarikat",
            "Registration Number"
        ]
    ),

    "EMAIL": combine_first_available(
        elsa,
        [
            "Emel",
            "Email"
        ]
    ),

    "PHONE": combine_first_available(
        elsa,
        [
            "Nombor telefon",
            "No Telefon",
            "Contact No"
        ]
    ),

    "STATE": combine_first_available(
        elsa,
        [
            "Negeri",
            "State"
        ]
    ),

    "SECTOR": combine_first_available(
        elsa,
        [
            "Sektor",
            "Sector",
            "Industry"
        ]
    ),

    "SME_CATEGORY": combine_first_available(
        elsa,
        [
            "Kategori SME",
            "SME Category"
        ]
    ),

    "ANNUAL_TURNOVER": combine_first_available(
        elsa,
        [
            "Perolehan Tahunan RM",
            "Perolehan Tahunan",
            "Annual Turnover"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "FULL_TIME_EMPLOYEES": combine_first_available(
        elsa,
        [
            "Berapakah jumlah pekerja yang bekerja secara sepenuh masa Sila nyatakan angka",
            "Jumlah Pekerja Sepenuh Masa",
            "Full Time Employees"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "HAS_BUSINESS_BANK_ACCOUNT": combine_first_available(
        elsa,
        [
            "Adakah anda mempunyai akaun bank berdaftar untuk syarikat anda",
            "Akaun Bank Perniagaan"
        ],
        default_value=np.nan
    ).map(convert_yes_no),

    "HAS_AUDITED_ACCOUNT": combine_first_available(
        elsa,
        [
            "Adakah anda mempunyai akaun perniagaan yang diaudit",
            "Akaun Diaudit"
        ],
        default_value=np.nan
    ).map(convert_yes_no),

    "SST_REGISTERED": combine_first_available(
        elsa,
        [
            "Adakah anda berdaftar dengan SST",
            "SST Registered"
        ],
        default_value=np.nan
    ).map(convert_yes_no),

    "HAS_BUSINESS_PLAN": combine_first_available(
        elsa,
        [
            "Adakah anda mempunyai Perancangan Perniagaan yang didokumentasikan untuk syarikat anda",
            "Business Plan"
        ],
        default_value=np.nan
    ).map(convert_yes_no),

    "ATTEMPTED_FINANCING": combine_first_available(
        elsa,
        [
            "Adakah anda pernah mencuba mendapatkan pembiayaan untuk perniagaan anda sejak dua tahun kebelakangan ini",
            "Attempted Financing"
        ],
        default_value=np.nan
    ).map(convert_yes_no),

    "SOURCE_SCORE": combine_first_available(
        elsa,
        [
            "SCORE",
            "TOTAL SCORE",
            "ELSA SCORE"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "ATTENDANCE": np.nan,
    "COLLECTION_RATE": np.nan,
    "TOTAL_COLLECTION": np.nan,
    "OUTSTANDING_BALANCE": np.nan,
    "AGEING": np.nan,
    "ECL": np.nan,
    "BUSINESS_NEED": "",
    "REGISTRATION_CATEGORY": "",
    "REFERRAL_SOURCE": ""
})

print(
    f"ELSA lead records: {len(elsa_leads):,}"
)


# ============================================================
# 8. PREPARE EPC LEADS
# ============================================================

print_section("PREPARING EPC LEADS")

epc_leads = pd.DataFrame({
    "LEAD_SOURCE": "EPC",
    "LEAD_DATE": "",

    "COMPANY_NAME": combine_first_available(
        epc,
        [
            "PENYEWA",
            "COMPANY NAME",
            "CUSTOMER NAME",
            "NAMA SYARIKAT"
        ]
    ),

    "REGISTRATION_NUMBER": combine_first_available(
        epc,
        [
            "REGISTRATION NUMBER",
            "NO PENDAFTARAN SYARIKAT"
        ]
    ),

    "EMAIL": combine_first_available(
        epc,
        [
            "EMAIL",
            "EMEL"
        ]
    ),

    "PHONE": combine_first_available(
        epc,
        [
            "CONTACT NO",
            "PHONE",
            "NO TELEFON"
        ]
    ),

    "STATE": combine_first_available(
        epc,
        [
            "STATE",
            "NEGERI"
        ]
    ),

    "SECTOR": combine_first_available(
        epc,
        [
            "INDUSTRY",
            "SECTOR",
            "SEKTOR"
        ]
    ),

    "SME_CATEGORY": "",

    "ANNUAL_TURNOVER": np.nan,
    "FULL_TIME_EMPLOYEES": np.nan,
    "HAS_BUSINESS_BANK_ACCOUNT": np.nan,
    "HAS_AUDITED_ACCOUNT": np.nan,
    "SST_REGISTERED": np.nan,
    "HAS_BUSINESS_PLAN": np.nan,
    "ATTEMPTED_FINANCING": np.nan,
    "SOURCE_SCORE": np.nan,
    "ATTENDANCE": np.nan,

    "COLLECTION_RATE": combine_first_available(
        epc,
        [
            "% TOTAL COLLECTION/DUE JUNE 2026",
            "% TIBAMASA VS KUTIPAN JUNE 2026",
            "COLLECTION RATE",
            "COLLECTION PERCENTAGE"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "TOTAL_COLLECTION": combine_first_available(
        epc,
        [
            "JUMLAH KUTIPAN JUNE 2026 RM",
            "TOTAL COLLECTION"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "OUTSTANDING_BALANCE": combine_first_available(
        epc,
        [
            "TMS BALANCE TOS JUNE 2026 RM",
            "OUTSTANDING BALANCE",
            "TMS BALANCE"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "AGEING": combine_first_available(
        epc,
        [
            "AGEING JUNE 2026",
            "AGEING",
            "AGING"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "ECL": combine_first_available(
        epc,
        [
            "ECL PROVISION FOR DOUBTFUL JUNE 2026",
            "ECL",
            "ECL PROVISION"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "BUSINESS_NEED": "",
    "REGISTRATION_CATEGORY": "",
    "REFERRAL_SOURCE": ""
})

print(
    f"EPC lead records: {len(epc_leads):,}"
)


# ============================================================
# 9. PREPARE JOMXCESS LEADS
# ============================================================

print_section("PREPARING JOMXCESS LEADS")

jomxcess_leads = pd.DataFrame({
    "LEAD_SOURCE": "JOMXCESS",

    "LEAD_DATE": combine_first_available(
        jomxcess,
        [
            "Timestamp",
            "Date",
            "Tarikh"
        ]
    ),

    "COMPANY_NAME": combine_first_available(
        jomxcess,
        [
            "Nama Syarikat",
            "Nama Syarikat 2",
            "Company Name"
        ]
    ),

    "REGISTRATION_NUMBER": combine_first_available(
        jomxcess,
        [
            "Nombor Pendaftaran",
            "No Pendaftaran Syarikat",
            "Registration Number"
        ]
    ),

    "EMAIL": combine_first_available(
        jomxcess,
        [
            "Email",
            "Emel"
        ]
    ),

    "PHONE": combine_first_available(
        jomxcess,
        [
            "No telefon pejabat",
            "No telefon bimbit pejabat",
            "No Telefon",
            "Contact No"
        ]
    ),

    "STATE": combine_first_available(
        jomxcess,
        [
            "Negeri",
            "State"
        ]
    ),

    "SECTOR": combine_first_available(
        jomxcess,
        [
            "Jenis Perniagaan Nature of Business",
            "Nature of Business",
            "Sector",
            "Industry"
        ]
    ),

    "SME_CATEGORY": combine_first_available(
        jomxcess,
        [
            "Kategori",
            "Kategori 2",
            "SME Category"
        ]
    ),

    "ANNUAL_TURNOVER": np.nan,
    "FULL_TIME_EMPLOYEES": np.nan,
    "HAS_BUSINESS_BANK_ACCOUNT": np.nan,
    "HAS_AUDITED_ACCOUNT": np.nan,
    "SST_REGISTERED": np.nan,
    "HAS_BUSINESS_PLAN": np.nan,
    "ATTEMPTED_FINANCING": np.nan,
    "SOURCE_SCORE": np.nan,

    "ATTENDANCE": combine_first_available(
        jomxcess,
        [
            "Jumlah kehadiran",
            "Jumlah kehadiran 2",
            "Attendance"
        ],
        default_value=np.nan
    ).map(convert_numeric),

    "COLLECTION_RATE": np.nan,
    "TOTAL_COLLECTION": np.nan,
    "OUTSTANDING_BALANCE": np.nan,
    "AGEING": np.nan,
    "ECL": np.nan,

    "BUSINESS_NEED": combine_first_available(
        jomxcess,
        [
            "Keperluan anda",
            "Business Need",
            "Requirement"
        ]
    ),

    "REGISTRATION_CATEGORY": combine_first_available(
        jomxcess,
        [
            "Kategori Pendaftaran anda",
            "Registration Category"
        ]
    ),

    "REFERRAL_SOURCE": combine_first_available(
        jomxcess,
        [
            "Sumber Rujukan",
            "Referral Source"
        ]
    )
})

print(
    f"JOMXCESS lead records: "
    f"{len(jomxcess_leads):,}"
)


# ============================================================
# 10. COMBINE LEADS
# ============================================================

print_section("COMBINING LEADS")

lead_data = pd.concat(
    [
        elsa_leads,
        epc_leads,
        jomxcess_leads
    ],
    ignore_index=True,
    sort=False
)


lead_data.insert(
    0,
    "LEAD_ID",
    [
        f"LEAD-{number:06d}"
        for number in range(
            1,
            len(lead_data) + 1
        )
    ]
)


lead_data["MATCH_NAME"] = (
    lead_data["COMPANY_NAME"]
    .map(normalise_company_name)
)

lead_data["MATCH_REGISTRATION"] = (
    lead_data["REGISTRATION_NUMBER"]
    .map(normalise_registration)
)

lead_data["MATCH_EMAIL"] = (
    lead_data["EMAIL"]
    .map(normalise_email)
)

lead_data["MATCH_PHONE"] = (
    lead_data["PHONE"]
    .map(normalise_phone)
)


# Remove records with no usable identifying information
lead_data = lead_data.loc[
    lead_data[
        [
            "MATCH_NAME",
            "MATCH_REGISTRATION",
            "MATCH_EMAIL",
            "MATCH_PHONE"
        ]
    ].ne("").any(axis=1)
].copy()


lead_data = lead_data.reset_index(drop=True)

print(
    f"Total usable lead records: "
    f"{len(lead_data):,}"
)


# ============================================================
# 11. MATCH LEADS TO EXISTING CUSTOMERS
# ============================================================

print_section("MATCHING LEADS TO CBRS AND MIS4")

lead_data["MATCHED_BY_REGISTRATION"] = (
    lead_data["MATCH_REGISTRATION"].ne("")
    & lead_data["MATCH_REGISTRATION"].isin(
        customer_registrations
    )
)

lead_data["MATCHED_BY_EMAIL"] = (
    lead_data["MATCH_EMAIL"].ne("")
    & lead_data["MATCH_EMAIL"].isin(
        customer_emails
    )
)

lead_data["MATCHED_BY_PHONE"] = (
    lead_data["MATCH_PHONE"].ne("")
    & lead_data["MATCH_PHONE"].isin(
        customer_phones
    )
)

lead_data["MATCHED_BY_NAME"] = (
    lead_data["MATCH_NAME"].ne("")
    & lead_data["MATCH_NAME"].isin(
        customer_names
    )
)


match_columns = [
    "MATCHED_BY_REGISTRATION",
    "MATCHED_BY_EMAIL",
    "MATCHED_BY_PHONE",
    "MATCHED_BY_NAME"
]


lead_data["MATCH_COUNT"] = (
    lead_data[match_columns]
    .astype(int)
    .sum(axis=1)
)


lead_data["CONVERTED"] = (
    lead_data["MATCH_COUNT"] > 0
).astype(int)


def get_match_method(row):
    methods = []

    if row["MATCHED_BY_REGISTRATION"]:
        methods.append("REGISTRATION")

    if row["MATCHED_BY_EMAIL"]:
        methods.append("EMAIL")

    if row["MATCHED_BY_PHONE"]:
        methods.append("PHONE")

    if row["MATCHED_BY_NAME"]:
        methods.append("COMPANY_NAME")

    if not methods:
        return "NO_MATCH"

    return ", ".join(methods)


lead_data["MATCH_METHOD"] = lead_data.apply(
    get_match_method,
    axis=1
)


print(
    "Matched to existing customers:",
    f"{lead_data['CONVERTED'].sum():,}"
)

print(
    "Not matched to existing customers:",
    f"{(lead_data['CONVERTED'] == 0).sum():,}"
)


# ============================================================
# 12. CALCULATE READINESS SCORE
# ============================================================

print_section("CALCULATING READINESS SCORES")


def calculate_readiness_score(row):
    earned_points = 0.0
    available_points = 0.0

    def add_binary(value, points):
        nonlocal earned_points
        nonlocal available_points

        if pd.notna(value):
            available_points += points

            if float(value) == 1:
                earned_points += points

    def add_numeric(value, points, upper_limit):
        nonlocal earned_points
        nonlocal available_points

        if pd.notna(value):
            numeric_value = max(
                float(value),
                0
            )

            available_points += points

            ratio = min(
                numeric_value / upper_limit,
                1
            )

            earned_points += ratio * points

    # ELSA indicators
    add_binary(
        row.get("HAS_BUSINESS_BANK_ACCOUNT"),
        12
    )

    add_binary(
        row.get("HAS_AUDITED_ACCOUNT"),
        12
    )

    add_binary(
        row.get("SST_REGISTERED"),
        8
    )

    add_binary(
        row.get("HAS_BUSINESS_PLAN"),
        8
    )

    add_binary(
        row.get("ATTEMPTED_FINANCING"),
        10
    )

    annual_turnover = row.get(
        "ANNUAL_TURNOVER"
    )

    if pd.notna(annual_turnover):
        available_points += 15

        turnover = max(
            float(annual_turnover),
            0
        )

        turnover_ratio = min(
            np.log1p(turnover)
            / np.log1p(50_000_000),
            1
        )

        earned_points += (
            turnover_ratio * 15
        )

    add_numeric(
        row.get("FULL_TIME_EMPLOYEES"),
        points=8,
        upper_limit=75
    )

    source_score = row.get(
        "SOURCE_SCORE"
    )

    if pd.notna(source_score):
        available_points += 12

        source_score = float(source_score)

        if source_score <= 5:
            source_ratio = source_score / 5
        elif source_score <= 100:
            source_ratio = source_score / 100
        else:
            source_ratio = 1

        source_ratio = min(
            max(source_ratio, 0),
            1
        )

        earned_points += (
            source_ratio * 12
        )

    # EPC indicators
    collection_rate = row.get(
        "COLLECTION_RATE"
    )

    if pd.notna(collection_rate):
        available_points += 30

        collection_rate = float(
            collection_rate
        )

        if collection_rate <= 1.5:
            collection_rate *= 100

        collection_ratio = min(
            max(collection_rate, 0),
            100
        ) / 100

        earned_points += (
            collection_ratio * 30
        )

    ageing = row.get("AGEING")

    if pd.notna(ageing):
        available_points += 20

        ageing = max(
            float(ageing),
            0
        )

        ageing_ratio = max(
            0,
            1 - min(ageing, 180) / 180
        )

        earned_points += (
            ageing_ratio * 20
        )

    ecl = row.get("ECL")

    if pd.notna(ecl):
        available_points += 10

        if float(ecl) <= 0:
            earned_points += 10

    # JOMXCESS indicators
    add_numeric(
        row.get("ATTENDANCE"),
        points=20,
        upper_limit=5
    )

    registration_category = str(
        row.get(
            "REGISTRATION_CATEGORY",
            ""
        )
    ).upper().strip()

    if registration_category:
        available_points += 10

        if "RSVP" in registration_category:
            earned_points += 10
        elif "WALK" in registration_category:
            earned_points += 5
        else:
            earned_points += 3

    business_need = str(
        row.get(
            "BUSINESS_NEED",
            ""
        )
    ).upper().strip()

    if business_need:
        available_points += 15

        financing_terms = [
            "PEMBIAYAAN",
            "FINANCING",
            "FINANCE",
            "LOAN"
        ]

        development_terms = [
            "PEMBANGUNAN",
            "DEVELOPMENT",
            "CAPACITY"
        ]

        if any(
            term in business_need
            for term in financing_terms
        ):
            earned_points += 15

        elif any(
            term in business_need
            for term in development_terms
        ):
            earned_points += 10

        else:
            earned_points += 5

    # Neutral score if no readiness information exists
    if available_points == 0:
        return 50.0

    final_score = (
        earned_points
        / available_points
        * 100
    )

    return round(
        min(
            max(final_score, 0),
            100
        ),
        2
    )


lead_data["READINESS_SCORE"] = (
    lead_data.apply(
        calculate_readiness_score,
        axis=1
    )
)


# ============================================================
# 13. PREPARE MODEL DATA
# ============================================================

print_section("PREPARING MODEL DATA")

model_features = [
    "LEAD_SOURCE",
    "STATE",
    "SECTOR",
    "SME_CATEGORY",
    "ANNUAL_TURNOVER",
    "FULL_TIME_EMPLOYEES",
    "HAS_BUSINESS_BANK_ACCOUNT",
    "HAS_AUDITED_ACCOUNT",
    "SST_REGISTERED",
    "HAS_BUSINESS_PLAN",
    "ATTEMPTED_FINANCING",
    "SOURCE_SCORE",
    "ATTENDANCE",
    "COLLECTION_RATE",
    "TOTAL_COLLECTION",
    "OUTSTANDING_BALANCE",
    "AGEING",
    "ECL",
    "BUSINESS_NEED",
    "REGISTRATION_CATEGORY",
    "REFERRAL_SOURCE",
    "READINESS_SCORE"
]


X = lead_data[
    model_features
].copy()

y = lead_data[
    "CONVERTED"
].copy()


categorical_features = [
    "LEAD_SOURCE",
    "STATE",
    "SECTOR",
    "SME_CATEGORY",
    "BUSINESS_NEED",
    "REGISTRATION_CATEGORY",
    "REFERRAL_SOURCE"
]


numeric_features = [
    column
    for column in model_features
    if column not in categorical_features
]


for column in categorical_features:
    X[column] = (
        X[column]
        .fillna("MISSING")
        .astype(str)
        .replace("", "MISSING")
    )


for column in numeric_features:
    X[column] = pd.to_numeric(
        X[column],
        errors="coerce"
    )

    median_value = X[column].median()

    if pd.isna(median_value):
        median_value = 0

    X[column] = X[column].fillna(
        median_value
    )


categorical_indices = [
    X.columns.get_loc(column)
    for column in categorical_features
]


# ============================================================
# 14. MODEL TRAINING
# ============================================================

print_section("TRAINING MODEL")

metrics = {
    "generated_at": (
        datetime.now().isoformat()
    ),
    "total_records": int(
        len(lead_data)
    ),
    "matched_records": int(
        lead_data["CONVERTED"].sum()
    ),
    "unmatched_records": int(
        (lead_data["CONVERTED"] == 0).sum()
    ),
    "model_trained": False,
    "warning": None
}


class_counts = y.value_counts()

can_train = (
    y.nunique() == 2
    and class_counts.min() >= 2
    and len(y) >= 10
)


feature_importance = pd.DataFrame({
    "FEATURE": ["READINESS_SCORE"],
    "IMPORTANCE": [100.0]
})


if can_train:

    groups = (
        lead_data["MATCH_NAME"]
        .replace("", np.nan)
    )

    fallback_groups = pd.Series(
        [
            f"ROW-{index}"
            for index in lead_data.index
        ],
        index=lead_data.index
    )

    groups = groups.fillna(
        fallback_groups
    )

    splitter = GroupShuffleSplit(
        n_splits=1,
        test_size=0.20,
        random_state=RANDOM_STATE
    )

    try:
        train_index, test_index = next(
            splitter.split(
                X,
                y,
                groups=groups
            )
        )

        X_train = X.iloc[
            train_index
        ].copy()

        X_test = X.iloc[
            test_index
        ].copy()

        y_train = y.iloc[
            train_index
        ].copy()

        y_test = y.iloc[
            test_index
        ].copy()

        if y_train.nunique() < 2:
            raise ValueError(
                "Training sample contains only one target class."
            )

        model = CatBoostClassifier(
            iterations=500,
            depth=6,
            learning_rate=0.03,
            loss_function="Logloss",
            eval_metric="AUC",
            auto_class_weights="Balanced",
            random_seed=RANDOM_STATE,
            verbose=False,
            allow_writing_files=False
        )

        model.fit(
            X_train,
            y_train,
            cat_features=categorical_indices
        )

        full_probabilities = (
            model.predict_proba(X)[:, 1]
        )

        lead_data[
            "CONVERSION_PROBABILITY"
        ] = (
            full_probabilities * 100
        )

        test_probabilities = (
            model.predict_proba(
                X_test
            )[:, 1]
        )

        test_predictions = (
            test_probabilities >= 0.50
        ).astype(int)

        metrics["model_trained"] = True
        metrics["training_records"] = int(
            len(train_index)
        )
        metrics["testing_records"] = int(
            len(test_index)
        )

        metrics["accuracy"] = float(
            accuracy_score(
                y_test,
                test_predictions
            )
        )

        metrics["precision"] = float(
            precision_score(
                y_test,
                test_predictions,
                zero_division=0
            )
        )

        metrics["recall"] = float(
            recall_score(
                y_test,
                test_predictions,
                zero_division=0
            )
        )

        metrics["f1_score"] = float(
            f1_score(
                y_test,
                test_predictions,
                zero_division=0
            )
        )

        metrics["confusion_matrix"] = (
            confusion_matrix(
                y_test,
                test_predictions,
                labels=[0, 1]
            ).tolist()
        )

        if y_test.nunique() == 2:
            metrics["roc_auc"] = float(
                roc_auc_score(
                    y_test,
                    test_probabilities
                )
            )

            metrics["pr_auc"] = float(
                average_precision_score(
                    y_test,
                    test_probabilities
                )
            )

        else:
            metrics["roc_auc"] = None
            metrics["pr_auc"] = None
            metrics["warning"] = (
                "The testing group contains only one "
                "target class. ROC-AUC and PR-AUC "
                "were not calculated."
            )

        feature_importance = pd.DataFrame({
            "FEATURE": X.columns,
            "IMPORTANCE": (
                model.get_feature_importance()
            )
        }).sort_values(
            "IMPORTANCE",
            ascending=False
        )

        model.save_model(
            MODEL_FILE
        )

        print(
            "CatBoost model trained successfully."
        )

    except Exception as error:
        metrics["model_trained"] = False
        metrics["warning"] = str(error)

        lead_data[
            "CONVERSION_PROBABILITY"
        ] = lead_data[
            "READINESS_SCORE"
        ]

        print(
            "Model training could not be completed."
        )

        print(
            "Using readiness score as the "
            "preliminary probability."
        )

        print("Reason:", error)

else:
    metrics["warning"] = (
        "The dataset does not have enough examples "
        "of both matched and unmatched leads to train "
        "a reliable binary classifier."
    )

    lead_data[
        "CONVERSION_PROBABILITY"
    ] = lead_data[
        "READINESS_SCORE"
    ]

    print(
        "Insufficient target classes for model training."
    )

    print(
        "Using readiness score as the "
        "preliminary probability."
    )


# ============================================================
# 15. FINAL LEAD SCORE
# ============================================================

print_section("CALCULATING FINAL LEAD SCORES")

lead_data["CONVERSION_PROBABILITY"] = (
    pd.to_numeric(
        lead_data["CONVERSION_PROBABILITY"],
        errors="coerce"
    )
    .fillna(50)
    .clip(0, 100)
    .round(2)
)


lead_data["FINAL_LEAD_SCORE"] = (
    (
        lead_data[
            "CONVERSION_PROBABILITY"
        ] * PROPENSITY_WEIGHT
    )
    +
    (
        lead_data[
            "READINESS_SCORE"
        ] * READINESS_WEIGHT
    )
).round(2)


def assign_priority(row):
    if row["CONVERTED"] == 1:
        return "EXISTING CUSTOMER"

    score = row["FINAL_LEAD_SCORE"]

    if score >= VERY_HOT_THRESHOLD:
        return "VERY HOT"

    if score >= HOT_THRESHOLD:
        return "HOT"

    if score >= WARM_THRESHOLD:
        return "WARM"

    return "COLD"


lead_data["LEAD_PRIORITY"] = (
    lead_data.apply(
        assign_priority,
        axis=1
    )
)


action_mapping = {
    "VERY HOT": (
        "Immediate relationship manager review"
    ),
    "HOT": (
        "Contact for qualification and financing discussion"
    ),
    "WARM": (
        "Continue engagement and collect more information"
    ),
    "COLD": (
        "Retain in lead pool for future nurturing"
    ),
    "EXISTING CUSTOMER": (
        "Exclude from new lead campaign"
    )
}


lead_data["RECOMMENDED_ACTION"] = (
    lead_data["LEAD_PRIORITY"]
    .map(action_mapping)
)


# ============================================================
# 16. DEDUPLICATE LEADS
# ============================================================

print_section("DEDUPLICATING LEADS")

lead_data["DEDUPLICATION_KEY"] = np.where(
    lead_data["MATCH_REGISTRATION"].ne(""),
    "REG|" + lead_data["MATCH_REGISTRATION"],
    np.where(
        lead_data["MATCH_EMAIL"].ne(""),
        "EMAIL|" + lead_data["MATCH_EMAIL"],
        np.where(
            lead_data["MATCH_PHONE"].ne(""),
            "PHONE|" + lead_data["MATCH_PHONE"],
            "NAME|" + lead_data["MATCH_NAME"]
        )
    )
)


ranked_data = (
    lead_data.sort_values(
        by=[
            "CONVERTED",
            "FINAL_LEAD_SCORE",
            "READINESS_SCORE"
        ],
        ascending=[
            True,
            False,
            False
        ]
    )
    .drop_duplicates(
        subset=[
            "LEAD_SOURCE",
            "DEDUPLICATION_KEY"
        ],
        keep="first"
    )
    .reset_index(drop=True)
)


new_leads = (
    ranked_data.loc[
        ranked_data["CONVERTED"] == 0
    ]
    .sort_values(
        by="FINAL_LEAD_SCORE",
        ascending=False
    )
    .reset_index(drop=True)
)


existing_customer_matches = (
    ranked_data.loc[
        ranked_data["CONVERTED"] == 1
    ]
    .sort_values(
        by=[
            "COMPANY_NAME",
            "LEAD_SOURCE"
        ]
    )
    .reset_index(drop=True)
)


new_leads.insert(
    0,
    "LEAD_RANK",
    range(
        1,
        len(new_leads) + 1
    )
)


print(
    f"Unique new leads: {len(new_leads):,}"
)

print(
    "Unique existing-customer matches:",
    f"{len(existing_customer_matches):,}"
)


# ============================================================
# 17. CREATE OUTPUT TABLES
# ============================================================

new_lead_columns = [
    "LEAD_RANK",
    "LEAD_ID",
    "LEAD_SOURCE",
    "COMPANY_NAME",
    "REGISTRATION_NUMBER",
    "EMAIL",
    "PHONE",
    "STATE",
    "SECTOR",
    "SME_CATEGORY",
    "BUSINESS_NEED",
    "ANNUAL_TURNOVER",
    "FULL_TIME_EMPLOYEES",
    "ATTENDANCE",
    "COLLECTION_RATE",
    "AGEING",
    "READINESS_SCORE",
    "CONVERSION_PROBABILITY",
    "FINAL_LEAD_SCORE",
    "LEAD_PRIORITY",
    "RECOMMENDED_ACTION"
]


new_lead_columns = [
    column
    for column in new_lead_columns
    if column in new_leads.columns
]


final_new_leads = new_leads[
    new_lead_columns
].copy()


existing_columns = [
    "LEAD_ID",
    "LEAD_SOURCE",
    "COMPANY_NAME",
    "REGISTRATION_NUMBER",
    "EMAIL",
    "PHONE",
    "MATCH_METHOD",
    "MATCH_COUNT",
    "READINESS_SCORE",
    "CONVERSION_PROBABILITY",
    "FINAL_LEAD_SCORE",
    "LEAD_PRIORITY",
    "RECOMMENDED_ACTION"
]


existing_columns = [
    column
    for column in existing_columns
    if column in existing_customer_matches.columns
]


final_existing_customers = (
    existing_customer_matches[
        existing_columns
    ].copy()
)


matching_detail_columns = [
    "LEAD_ID",
    "LEAD_SOURCE",
    "COMPANY_NAME",
    "REGISTRATION_NUMBER",
    "EMAIL",
    "PHONE",
    "MATCH_NAME",
    "MATCH_REGISTRATION",
    "MATCH_EMAIL",
    "MATCH_PHONE",
    "MATCHED_BY_REGISTRATION",
    "MATCHED_BY_EMAIL",
    "MATCHED_BY_PHONE",
    "MATCHED_BY_NAME",
    "MATCH_COUNT",
    "MATCH_METHOD",
    "CONVERTED"
]


matching_details = lead_data[
    matching_detail_columns
].copy()


priority_summary = (
    final_new_leads.groupby(
        [
            "LEAD_SOURCE",
            "LEAD_PRIORITY"
        ],
        dropna=False
    )
    .size()
    .reset_index(
        name="TOTAL_LEADS"
    )
)


if len(final_new_leads) > 0:
    source_summary = (
        final_new_leads.groupby(
            "LEAD_SOURCE",
            dropna=False
        )
        .agg(
            TOTAL_LEADS=(
                "LEAD_ID",
                "count"
            ),
            AVERAGE_READINESS_SCORE=(
                "READINESS_SCORE",
                "mean"
            ),
            AVERAGE_CONVERSION_PROBABILITY=(
                "CONVERSION_PROBABILITY",
                "mean"
            ),
            AVERAGE_FINAL_LEAD_SCORE=(
                "FINAL_LEAD_SCORE",
                "mean"
            )
        )
        .reset_index()
    )

    score_columns = [
        "AVERAGE_READINESS_SCORE",
        "AVERAGE_CONVERSION_PROBABILITY",
        "AVERAGE_FINAL_LEAD_SCORE"
    ]

    source_summary[
        score_columns
    ] = source_summary[
        score_columns
    ].round(2)

else:
    source_summary = pd.DataFrame(
        columns=[
            "LEAD_SOURCE",
            "TOTAL_LEADS",
            "AVERAGE_READINESS_SCORE",
            "AVERAGE_CONVERSION_PROBABILITY",
            "AVERAGE_FINAL_LEAD_SCORE"
        ]
    )


# ============================================================
# 18. EXPORT EXCEL
# ============================================================

print_section("EXPORTING EXCEL")

with pd.ExcelWriter(
    OUTPUT_EXCEL,
    engine="openpyxl"
) as writer:

    final_new_leads.to_excel(
        writer,
        sheet_name="New Lead Rankings",
        index=False
    )

    final_existing_customers.to_excel(
        writer,
        sheet_name="Existing Customers",
        index=False
    )

    priority_summary.to_excel(
        writer,
        sheet_name="Priority Summary",
        index=False
    )

    source_summary.to_excel(
        writer,
        sheet_name="Source Summary",
        index=False
    )

    feature_importance.to_excel(
        writer,
        sheet_name="Feature Importance",
        index=False
    )

    matching_details.to_excel(
        writer,
        sheet_name="Matching Details",
        index=False
    )

    customer_master.to_excel(
        writer,
        sheet_name="Customer Master",
        index=False
    )


# ============================================================
# 19. FORMAT EXCEL
# ============================================================

print_section("FORMATTING EXCEL")

workbook = load_workbook(
    OUTPUT_EXCEL
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)


header_font = Font(
    color="FFFFFF",
    bold=True
)


priority_colours = {
    "VERY HOT": "C6EFCE",
    "HOT": "FFEB9C",
    "WARM": "FCE4D6",
    "COLD": "E7E6E6",
    "EXISTING CUSTOMER": "D9EAF7"
}


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    # Correct header formatting loop
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Adjust column widths
    for column_cells in worksheet.columns:
        max_length = 0

        column_number = (
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                cell_length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    cell_length
                )

        column_letter = get_column_letter(
            column_number
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            45
        )

    # Identify columns from header row
    header_map = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    priority_column = header_map.get(
        "LEAD_PRIORITY"
    )

    # Apply priority colours
    if priority_column is not None:

        for row_number in range(
            2,
            worksheet.max_row + 1
        ):
            priority_cell = worksheet.cell(
                row=row_number,
                column=priority_column
            )

            priority_value = (
                priority_cell.value
            )

            if priority_value in priority_colours:
                priority_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=priority_colours[
                        priority_value
                    ]
                )

                priority_cell.font = Font(
                    bold=True
                )


workbook.save(
    OUTPUT_EXCEL
)

print(
    f"Excel created: {OUTPUT_EXCEL}"
)


# ============================================================
# 20. EXPORT JSON
# ============================================================

print_section("EXPORTING JSON")

json_dataframe = final_new_leads.copy()


# Convert pandas missing values into Python None
json_dataframe = json_dataframe.replace(
    {
        np.nan: None,
        np.inf: None,
        -np.inf: None,
        pd.NaT: None
    }
)


json_records = json_dataframe.to_dict(
    orient="records"
)


with open(
    OUTPUT_JSON,
    "w",
    encoding="utf-8"
) as json_file:

    json.dump(
        json_records,
        json_file,
        ensure_ascii=False,
        indent=4,
        default=str
    )


print(
    f"JSON created: {OUTPUT_JSON}"
)


# ============================================================
# 21. EXPORT MODEL METRICS
# ============================================================

with open(
    METRICS_JSON,
    "w",
    encoding="utf-8"
) as metrics_file:

    json.dump(
        metrics,
        metrics_file,
        ensure_ascii=False,
        indent=4,
        default=str
    )


print(
    f"Metrics created: {METRICS_JSON}"
)


# ============================================================
# 22. FINAL SUMMARY
# ============================================================

print_section("PROCESS COMPLETED")

print("\nGenerated files:")

print(
    f"1. {OUTPUT_EXCEL}"
)

print(
    f"2. {OUTPUT_JSON}"
)

print(
    f"3. {METRICS_JSON}"
)

if metrics["model_trained"]:
    print(
        f"4. {MODEL_FILE}"
    )


print("\nResult summary:")

print(
    f"Total usable lead records: "
    f"{len(lead_data):,}"
)

print(
    f"Matched to existing customers: "
    f"{lead_data['CONVERTED'].sum():,}"
)

print(
    f"Unmatched lead records: "
    f"{(lead_data['CONVERTED'] == 0).sum():,}"
)

print(
    f"Unique new leads ranked: "
    f"{len(final_new_leads):,}"
)


if len(final_new_leads) > 0:

    display_columns = [
        "LEAD_RANK",
        "LEAD_SOURCE",
        "COMPANY_NAME",
        "READINESS_SCORE",
        "CONVERSION_PROBABILITY",
        "FINAL_LEAD_SCORE",
        "LEAD_PRIORITY"
    ]

    display_columns = [
        column
        for column in display_columns
        if column in final_new_leads.columns
    ]

    print("\nTop 20 leads:")

    print(
        final_new_leads[
            display_columns
        ].head(20).to_string(
            index=False
        )
    )

else:
    print(
        "\nNo new leads remained after matching "
        "and deduplication."
    )


if metrics.get("warning"):
    print(
        "\nModel warning:"
    )

    print(
        metrics["warning"]
    )


print(
    "\nImportant modelling note:"
)

print(
    "The current CONVERTED label means the lead matches "
    "a record in CBRS or MIS4. Before production use, "
    "verify that the lead date occurred before the "
    "customer relationship began."
)